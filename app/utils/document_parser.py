"""
文档解析器模块，用于解析不同类型的文档并提取文本内容
"""
import os
import PyPDF2
import pdfplumber
import docx
import logging
import uuid
import shutil
import hashlib
import mammoth
from datetime import datetime
from abc import ABC, abstractmethod
from app.core.config import UPLOADS_DIR, ALLOWED_EXTENSIONS, DOCUMENT_IMAGES_DIR, DOCUMENT_IMAGES_URL, ENV
from PIL import Image
from io import BytesIO

class DocumentParser(ABC):
    """文档解析器抽象基类"""

    @abstractmethod
    def parse(self, file_path):
        """
        解析文档内容

        Args:
            file_path: 文档路径

        Returns:
            dict: 包含解析后的文本内容和元数据
        """
        pass

    @staticmethod
    def get_parser(file_path):
        """
        根据文件类型获取合适的解析器
        
        Args:
            file_path: 文件路径
            
        Returns:
            DocumentParser: 适合的解析器实例
        """
        ext = os.path.splitext(file_path)[1].lower()
        
        if ext == '.pdf':
            return PdfParser()
        elif ext in ('.docx', '.doc'):
            return DocxParser()
        elif ext == '.txt':
            return TxtParser()
        else:
            raise ValueError(f"不支持的文件类型: {ext}")


class PdfParser(DocumentParser):
    """PDF文档解析器"""
    
    def parse(self, file_path, doc_id=None):
        """
        解析PDF文档内容
        
        Args:
            file_path: 文档路径
            doc_id: 文档ID（如果有）
            
        Returns:
            dict: 包含文本内容和元数据的字典
        """
        text = ""
        metadata = {}
        
        try:
            # 使用PyPDF2提取元数据
            with open(file_path, 'rb') as f:
                pdf_reader = PyPDF2.PdfReader(f)
                
                # 提取元数据
                info = pdf_reader.metadata
                if info:
                    metadata = {
                        'title': info.title or '',
                        'author': info.author or '',
                        'subject': info.subject or '',
                        'creator': info.creator or '',
                        'producer': info.producer or '',
                        'pages': len(pdf_reader.pages)
                    }
                else:
                    metadata = {'pages': len(pdf_reader.pages)}
            
            # 使用pdfplumber提取文本（更准确）
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text() or ""
                    text += page_text + "\n\n"
                
        except Exception as e:
            logging.error(f"解析PDF文件失败: {e}")
            raise
            
        return {
            'text': text.strip(),
            'html_content': '',  # PDF暂不支持HTML格式
            'metadata': metadata
        }


class DocxParser(DocumentParser):
    """DOCX文档解析器"""
    
    def _get_image_directory(self, file_path, doc_id):
        """获取文档图片目录"""
        # 使用配置的文档图片目录
        # 按年月组织图片，便于管理
        today = datetime.now()
        year_month = today.strftime('%Y%m')
        
        if doc_id:
            # 使用文档ID创建目录
            image_dir = os.path.join(DOCUMENT_IMAGES_DIR, year_month, str(doc_id))
            # URL路径使用环境特定的前缀
            rel_path = f'{DOCUMENT_IMAGES_URL}/{year_month}/{doc_id}'
        else:
            # 如果没有文档ID，使用文件名哈希作为目录
            doc_filename = os.path.basename(file_path)
            file_hash = hashlib.md5(doc_filename.encode()).hexdigest()[:10]
            image_dir = os.path.join(DOCUMENT_IMAGES_DIR, year_month, file_hash)
            rel_path = f'{DOCUMENT_IMAGES_URL}/{year_month}/{file_hash}'
        
        # 确保目录存在 - 增强的目录创建逻辑
        if not os.path.exists(image_dir):
            try:
                # 先确保父目录存在
                parent_dir = os.path.dirname(image_dir)
                if not os.path.exists(parent_dir):
                    os.makedirs(parent_dir, exist_ok=True)
                    logging.info(f"创建父目录: {parent_dir}")
                
                # 然后创建图片目录
                os.makedirs(image_dir, exist_ok=True)
                logging.info(f"创建图片目录: {image_dir}")
                
                # 验证目录是否真的创建成功
                if not os.path.exists(image_dir):
                    raise IOError(f"目录创建失败: {image_dir}")
                
                # 验证目录是否可写
                test_file = os.path.join(image_dir, '.test')
                try:
                    with open(test_file, 'w') as f:
                        f.write('test')
                    os.remove(test_file)
                    logging.info(f"目录权限检查通过: {image_dir}")
                except Exception as e:
                    raise IOError(f"目录不可写: {image_dir}, 错误: {e}")
                
            except Exception as e:
                logging.error(f"创建图片目录失败: {e}, 路径: {image_dir}")
                # 尝试使用系统临时目录作为备用
                import tempfile
                temp_dir = os.path.join(tempfile.gettempdir(), 'docradar_images', str(doc_id) if doc_id else file_hash)
                try:
                    os.makedirs(temp_dir, exist_ok=True)
                    logging.warning(f"使用临时目录作为备用: {temp_dir}")
                    image_dir = temp_dir
                    rel_path = f'/static/temp_images/{str(doc_id) if doc_id else file_hash}'
                except Exception as temp_err:
                    logging.error(f"创建临时目录也失败: {temp_err}")
                    # 最后的备用方案，使用static目录
                    static_dir = os.path.join('app', 'static', 'images', 'docs', str(doc_id) if doc_id else file_hash)
                    os.makedirs(static_dir, exist_ok=True)
                    logging.warning(f"使用静态目录作为最后手段: {static_dir}")
                    image_dir = static_dir
                    rel_path = f'/static/images/docs/{str(doc_id) if doc_id else file_hash}'
        else:
            logging.info(f"图片目录已存在: {image_dir}")
        
        return image_dir, rel_path
    
    def _extract_images(self, file_path, doc_id=None):
        """提取DOCX文档中的图片，包括尺寸信息"""
        try:
            # 创建图片存储目录
            image_dir, rel_path = self._get_image_directory(file_path, doc_id)
            
            # 再次验证目录是否存在且可写
            if not os.path.exists(image_dir):
                logging.error(f"图片目录不存在，尝试重新创建: {image_dir}")
                os.makedirs(image_dir, exist_ok=True)
            
            # 验证目录可写性
            try:
                test_file = os.path.join(image_dir, '.write_test')
                with open(test_file, 'w') as f:
                    f.write('test')
                os.remove(test_file)
            except Exception as e:
                logging.error(f"图片目录不可写，尝试修复权限: {e}")
                # 尝试修复权限
                try:
                    import stat
                    os.chmod(image_dir, stat.S_IRWXU | stat.S_IRWXG | stat.S_IRWXO)  # 777权限
                except Exception as chmod_err:
                    logging.error(f"修复权限失败: {chmod_err}")
                    # 如果修复失败，尝试使用临时目录
                    import tempfile
                    temp_dir = os.path.join(tempfile.gettempdir(), 'docradar_temp_images')
                    os.makedirs(temp_dir, exist_ok=True)
                    image_dir = temp_dir
                    rel_path = '/static/temp_images'
                    logging.warning(f"使用临时目录作为备用: {image_dir}")
            
            # 打开docx文件
            from zipfile import ZipFile
            z = ZipFile(file_path)
            
            # 提取所有图片文件
            image_mappings = {}
            image_count = 0
            
            for zinfo in z.filelist:
                if zinfo.filename.startswith('word/media/'):
                    try:
                        # 获取原始文件名
                        image_name = os.path.basename(zinfo.filename)
                        # 构建新的图片路径和URL
                        image_path = os.path.join(image_dir, image_name)
                        image_url = f"{rel_path}/{image_name}"
                        
                        # 提取图片文件
                        image_data = z.read(zinfo.filename)
                        
                        # 确保目录存在
                        img_dir = os.path.dirname(image_path)
                        if not os.path.exists(img_dir):
                            os.makedirs(img_dir, exist_ok=True)
                        
                        # 写入图片文件
                        with open(image_path, 'wb') as f:
                            f.write(image_data)
                        
                        # 验证文件是否成功写入
                        if not os.path.exists(image_path) or os.path.getsize(image_path) == 0:
                            raise IOError(f"图片文件写入失败: {image_path}")
                        
                        # 获取图片尺寸
                        width = None
                        height = None
                        try:
                            with Image.open(BytesIO(image_data)) as img:
                                width, height = img.size
                                logging.info(f"图片尺寸: {width}x{height}")
                        except Exception as e:
                            logging.warning(f"无法获取图片尺寸: {e}")
                        
                        # 记录图片映射关系
                        image_mappings[image_name] = {
                            'path': image_path,
                            'url': image_url,
                            'name': image_name,
                            'width': width,
                            'height': height
                        }
                        image_count += 1
                        
                    except Exception as img_err:
                        logging.error(f"处理图片失败: {zinfo.filename}, 错误: {img_err}")
                        continue
            
            logging.info(f"从DOCX文档提取了 {image_count} 张图片")
            
            # 如果没有提取到任何图片，可能是文档格式问题
            if image_count == 0:
                logging.warning("未从文档中提取到任何图片，可能是文档格式问题")
            
            return image_mappings
            
        except Exception as e:
            logging.error(f"提取DOCX图片失败: {e}")
            # 返回空映射，确保不会因为图片提取失败而中断整个文档处理
            return {}
    
    def _convert_docx_to_html_mammoth(self, file_path, doc_id=None):
        """
        使用mammoth转换DOCX为HTML
        
        Args:
            file_path: DOCX文件路径
            doc_id: 文档ID
            
        Returns:
            str: 转换后的HTML内容
        """
        try:
            # 先提取图片，获取图片映射关系
            image_mappings = self._extract_images(file_path, doc_id)
            
            # 创建图片处理器，将文档中的图片引用替换为静态URL
            def image_handler(image):
                try:
                    if image.alt_text and image.alt_text in image_mappings:
                        img_info = image_mappings[image.alt_text]
                        result = {"src": img_info['url']}
                        if 'width' in img_info and img_info['width']:
                            result['width'] = str(img_info['width'])
                        if 'height' in img_info and img_info['height']:
                            result['height'] = str(img_info['height'])
                        return result
                    elif image.content_type and f"image.{image.content_type.split('/')[1]}" in image_mappings:
                        image_key = f"image.{image.content_type.split('/')[1]}"
                        img_info = image_mappings[image_key]
                        result = {"src": img_info['url']}
                        if 'width' in img_info and img_info['width']:
                            result['width'] = str(img_info['width'])
                        if 'height' in img_info and img_info['height']:
                            result['height'] = str(img_info['height'])
                        return result
                    else:
                        # 使用默认方式处理图片
                        try:
                            with image.open() as image_bytes:
                                # 生成唯一文件名
                                image_ext = 'png'  # 默认扩展名
                                if image.content_type and '/' in image.content_type:
                                    image_ext = image.content_type.split('/')[1]
                                image_name = f"img_{uuid.uuid4().hex}.{image_ext}"
                                
                                # 获取图片目录，确保存在
                                image_dir, rel_path = self._get_image_directory(file_path, doc_id)
                                if not os.path.exists(image_dir):
                                    os.makedirs(image_dir, exist_ok=True)
                                    logging.info(f"为内嵌图片创建目录: {image_dir}")
                                
                                image_path = os.path.join(image_dir, image_name)
                                image_url = f"{rel_path}/{image_name}"
                                
                                # 读取图片数据
                                try:
                                    image_data = image_bytes.read()
                                    if not image_data:
                                        raise ValueError("读取到的图片数据为空")
                                except Exception as read_err:
                                    logging.error(f"读取图片数据失败: {read_err}")
                                    # 返回空图片占位符
                                    return {"src": "/static/images/image_placeholder.png", "alt": "图片加载失败"}
                                
                                # 获取图片尺寸
                                width = None
                                height = None
                                try:
                                    with Image.open(BytesIO(image_data)) as img:
                                        width, height = img.size
                                        logging.info(f"图片尺寸: {width}x{height}")
                                except Exception as e:
                                    logging.warning(f"无法获取图片尺寸: {e}")
                                
                                # 保存图片
                                try:
                                    with open(image_path, "wb") as f:
                                        f.write(image_data)
                                    
                                    # 验证图片是否成功保存
                                    if not os.path.exists(image_path) or os.path.getsize(image_path) == 0:
                                        raise IOError(f"图片保存失败: {image_path}")
                                        
                                except Exception as save_err:
                                    logging.error(f"保存图片失败: {save_err}")
                                    # 尝试使用临时目录
                                    import tempfile
                                    temp_dir = os.path.join(tempfile.gettempdir(), 'docradar_images')
                                    os.makedirs(temp_dir, exist_ok=True)
                                    image_path = os.path.join(temp_dir, image_name)
                                    
                                    try:
                                        with open(image_path, "wb") as f:
                                            f.write(image_data)
                                        image_url = f"/static/temp_images/{image_name}"
                                    except Exception as temp_save_err:
                                        logging.error(f"保存到临时目录也失败: {temp_save_err}")
                                        # 最后的备用方案：返回空图片占位符
                                        return {"src": "/static/images/image_placeholder.png", "alt": "图片加载失败"}
                                
                                # 返回图片URL和尺寸信息
                                result = {"src": image_url}
                                if width and height:
                                    result['width'] = str(width)
                                    result['height'] = str(height)
                                return result
                                
                        except Exception as e:
                            logging.error(f"处理图片时发生错误: {e}")
                            # 返回空图片占位符
                            return {"src": "/static/images/image_placeholder.png", "alt": "图片加载失败"}
                except Exception as outer_e:
                    logging.error(f"图片处理器外层错误: {outer_e}")
                    # 最后的备用方案
                    return {"src": "/static/images/image_placeholder.png", "alt": "图片处理错误"}
            
            # 配置样式映射
            style_map = """
                p[style-name='Heading 1'] => h1:fresh
                p[style-name='Heading 2'] => h2:fresh
                p[style-name='Heading 3'] => h3:fresh
                p[style-name='Heading 4'] => h4:fresh
                p[style-name='Heading 5'] => h5:fresh
                p[style-name='Heading 6'] => h6:fresh
                r[style-name='Strong'] => strong
                r[style-name='Emphasis'] => em
                p[style-name='Quote'] => blockquote
                p[style-name='Intense Quote'] => blockquote.intense
                r[style-name='Underline'] => u
                p[style-name='List Paragraph'] => li
                table => table.table.table-bordered
            """
            
            # 执行转换 - 直接传递关键字参数而不是options字典
            result = mammoth.convert_to_html(
                file_path, 
                convert_image=mammoth.images.img_element(image_handler),
                style_map=style_map
            )
            html = result.value
            
            # 记录警告信息
            for message in result.messages:
                if message.type == "warning":
                    logging.warning(f"Mammoth警告: {message}")
            
            # 后处理HTML，为图片添加容器
            import re
            def wrap_images_with_container(html_content):
                # 匹配HTML中的img标签
                img_pattern = r'<img([^>]*)>'
                
                def replace_img(match):
                    img_tag = match.group(0)
                    # 将图片包装在容器中
                    return f'<div class="image-container">{img_tag}</div>'
                
                # 替换所有图片标签
                processed_html = re.sub(img_pattern, replace_img, html_content)
                return processed_html
            
            # 添加表格处理函数，将表格包装在可滚动容器中
            def wrap_tables_with_container(html_content):
                # 匹配HTML中的table标签
                table_pattern = r'<table([^>]*)>(.*?)</table>'
                
                def replace_table(match):
                    table_tag = match.group(0)
                    # 将表格包装在容器中
                    return f'<div class="table-container">{table_tag}</div>'
                
                # 替换所有表格标签
                processed_html = re.sub(table_pattern, replace_table, html_content, flags=re.DOTALL)
                return processed_html
            
            # 应用图片容器处理
            html = wrap_images_with_container(html)
            
            # 应用表格容器处理
            html = wrap_tables_with_container(html)
            
            # 添加兜底展示：如果有提取的图片但在HTML中找不到，在底部显示
            missing_images = []
            for img_name, img_info in image_mappings.items():
                if img_info['url'] not in html:
                    missing_images.append(img_info)
            
            if missing_images:
                html += '<div class="document-images-fallback"><h3>文档中的其他图片</h3><div class="image-gallery">'
                for img_info in missing_images:
                    width_attr = f' width="{img_info["width"]}"' if 'width' in img_info and img_info['width'] else ''
                    height_attr = f' height="{img_info["height"]}"' if 'height' in img_info and img_info['height'] else ''
                    html += f'<div class="image-item"><div class="image-container"><img src="{img_info["url"]}"{width_attr}{height_attr} alt="文档图片"></div><div class="image-caption">文档图片</div></div>'
                html += '</div></div>'
            
            # 添加基本样式
            styled_html = f"""
            <div class="document-content docx-document">
                <style>
                    .docx-document img {{ max-width: 100%; height: auto; }}
                    .document-images-fallback {{ margin-top: 2rem; padding-top: 1rem; border-top: 1px solid #dee2e6; }}
                    .image-gallery {{ display: flex; flex-wrap: wrap; gap: 1rem; margin-top: 1rem; }}
                    .image-item {{ width: auto; margin-bottom: 1rem; }}
                    .image-caption {{ font-size: 0.8rem; color: #6c757d; margin-top: 0.5rem; text-align: center; }}
                    .image-container {{ margin: 1rem 0; text-align: center; }}
                    .table-container {{ width: 100%; overflow-x: auto; margin-bottom: 1rem; }}
                </style>
                {html}
            </div>
            """
            
            return styled_html
            
        except Exception as e:
            logging.error(f"转换DOCX为HTML失败: {e}")
            # 转换失败时返回简单格式
            return f"<div class='error'>文档格式转换失败: {str(e)}</div>"
    
    def parse(self, file_path, doc_id=None):
        """解析DOCX文档内容，支持富文本格式和图片
        
        Args:
            file_path: 文档路径
            doc_id: 文档ID（如果有）
            
        Returns:
            dict: 包含文本内容、HTML内容和元数据的字典
        """
        text = ""
        html_content = ""
        metadata = {}
        
        try:
            doc = docx.Document(file_path)
            
            # 提取元数据
            core_properties = doc.core_properties
            metadata = {
                'title': core_properties.title or '',
                'author': core_properties.author or '',
                'subject': core_properties.subject or '',
                'created': str(core_properties.created) if core_properties.created else '',
                'modified': str(core_properties.modified) if core_properties.modified else '',
                'paragraphs': len(doc.paragraphs)
            }
            
            # 提取文本（用于搜索）
            for para in doc.paragraphs:
                text += para.text + "\n"
            
            # 使用mammoth生成高保真HTML
            html_content = self._convert_docx_to_html_mammoth(file_path, doc_id)
            
            # 保存图片目录信息，用于后续文档删除时清理
            image_dir, _ = self._get_image_directory(file_path, doc_id)
            metadata['image_directory'] = image_dir
                
        except Exception as e:
            logging.error(f"解析DOCX文件失败: {e}")
            raise
            
        return {
            'text': text.strip(),        # 纯文本用于搜索
            'html_content': html_content,  # HTML内容用于显示
            'metadata': metadata
        }


class TxtParser(DocumentParser):
    """TXT文档解析器"""
    
    def parse(self, file_path, doc_id=None):
        """
        解析TXT文档内容
        
        Args:
            file_path: 文档路径
            doc_id: 文档ID（如果有）
            
        Returns:
            dict: 包含文本内容和元数据的字典
        """
        text = ""
        
        try:
            # 检测编码
            encodings = ['utf-8', 'gbk', 'gb2312', 'latin-1']
            
            for encoding in encodings:
                try:
                    with open(file_path, 'r', encoding=encoding) as f:
                        text = f.read()
                    break
                except UnicodeDecodeError:
                    continue
            
            if not text:
                # 所有编码都失败，使用二进制模式读取
                with open(file_path, 'rb') as f:
                    text = f.read().decode('utf-8', errors='replace')
                    
            # 处理基本元数据
            file_stat = os.stat(file_path)
            metadata = {
                'size': file_stat.st_size,
                'created': file_stat.st_ctime,
                'modified': file_stat.st_mtime,
                'lines': text.count('\n') + 1,
                'words': len(text.split())
            }
            
            # 创建简单的HTML表示
            html_content = f"<pre>{text}</pre>"
                
        except Exception as e:
            logging.error(f"解析TXT文件失败: {e}")
            raise
            
        return {
            'text': text.strip(),
            'html_content': html_content,
            'metadata': metadata
        } 